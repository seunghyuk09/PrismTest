# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as gcl
from openpyxl.chart import LineChart, BarChart, Reference

SRC = '/root/.claude/uploads/2755f4b7-a3d7-5284-9153-97d4b1a5fed5/ecaa1d6d-U20_Plastic_Prism______From_ALT_2026.xlsx'
OUT = '/tmp/claude-0/-home-user-PrismTest/2755f4b7-a3d7-5284-9153-97d4b1a5fed5/scratchpad/xls/U20_프리즘_불량분석_2024-2026.xlsx'

DEFECTS = ['흰색 이물','검정색 이물','스크래치','Diffuser 얼룩','Dig','파손','내부 기포','내부 백테','표면 얼룩','페인트 얼룩']

# ---------- 1. 원자료 추출 ----------
src = openpyxl.load_workbook(SRC, data_only=True)
lots = []
for name in src.sheetnames:
    if name == 'Summary':
        continue
    ws = src[name]
    good = ws.cell(row=6, column=3).value or 0
    counts = [ws.cell(row=6, column=4+i).value or 0 for i in range(10)]
    total = ws.cell(row=8, column=3).value or 0
    date = name.split('_')[0]
    label = name.replace('_1',' (1차)').replace('_2',' (2차)')
    ym = date[:7]
    lots.append(dict(label=label, date=date, ym=ym, good=good, counts=counts, total=total))

for L in lots:
    assert L['good'] + sum(L['counts']) == L['total'], L['label']

months = []
for L in lots:
    if L['ym'] not in months:
        months.append(L['ym'])

# ---------- 서식 ----------
FONT = '맑은 고딕'
def F(sz=10, b=False, color='FF1F2328', it=False):
    return Font(name=FONT, size=sz, bold=b, color=color, italic=it)
HDR_FILL   = PatternFill('solid', fgColor='FF1F3A54')
SUB_FILL   = PatternFill('solid', fgColor='FFDCE4EC')
TOT_FILL   = PatternFill('solid', fgColor='FFF0F2F4')
KEY_FILL   = PatternFill('solid', fgColor='FFFFF3CD')
IN_FILL    = PatternFill('solid', fgColor='FFF7FAFD')
thin = Side(style='thin', color='FFB8C2CC')
med  = Side(style='medium', color='FF1F3A54')
BOX  = Border(left=thin, right=thin, top=thin, bottom=thin)
N0, N2, P2 = '#,##0', '#,##0.00', '0.00%'
C, L_, R = Alignment('center','center'), Alignment('left','center'), Alignment('right','center')

def head(ws, row, labels, start=1, h=30):
    for i, t in enumerate(labels):
        c = ws.cell(row=row, column=start+i, value=t)
        c.font = F(9.5, True, 'FFFFFFFF'); c.fill = HDR_FILL
        c.alignment = Alignment('center','center', wrap_text=True); c.border = BOX
    ws.row_dimensions[row].height = h

def title(ws, text, sub, ncol):
    ws['A1'] = text; ws['A1'].font = F(15, True, 'FF1F3A54')
    ws['A2'] = sub;  ws['A2'].font = F(9, color='FF5A6672', it=True)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
    ws.row_dimensions[1].height = 22

def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w

wb = openpyxl.Workbook()

# =========================================================
# 원본데이터  (다른 시트가 참조하는 단일 원천)
# =========================================================
raw = wb.active; raw.title = '원본데이터'
title(raw, 'U20 플라스틱 프리즘 — 입고검사 원자료',
      '출처: U20 Plastic Prism 검사현황 (From ALT 2026) 시트별 「소 계」 행. 회색 배경 = 원자료 입력값, 그 외 = 수식.', 18)
HDRS = ['No','납품일','연월','입고수량','양품','불량 계'] + DEFECTS + ['원자료 합계','검증']
head(raw, 4, HDRS)
R0 = 5
for i, L in enumerate(lots):
    r = R0 + i
    raw.cell(row=r, column=1, value=i+1).alignment = C
    raw.cell(row=r, column=2, value=L['label']).alignment = L_
    raw.cell(row=r, column=3, value=L['ym']).alignment = C
    raw.cell(row=r, column=4, value=f'=E{r}+F{r}')                     # 입고 = 양품+불량
    raw.cell(row=r, column=5, value=L['good'])                          # 입력
    raw.cell(row=r, column=6, value=f'=SUM(G{r}:P{r})')                 # 불량계
    for j, v in enumerate(L['counts']):
        raw.cell(row=r, column=7+j, value=v)                            # 입력
    raw.cell(row=r, column=17, value=L['total'])                        # 원자료 합계(입력)
    raw.cell(row=r, column=18, value=f'=D{r}-Q{r}')                     # 검증
    for cix in range(1, 19):
        c = raw.cell(row=r, column=cix); c.font = F(10); c.border = BOX
        if cix >= 4: c.number_format = N0; c.alignment = R
        if cix == 5 or 7 <= cix <= 16 or cix == 17: c.fill = IN_FILL
TR = R0 + len(lots)
raw.cell(row=TR, column=1, value='합계').alignment = C
raw.merge_cells(start_row=TR, start_column=1, end_row=TR, end_column=3)
for cix in range(4, 18):
    col = gcl(cix)
    raw.cell(row=TR, column=cix, value=f'=SUM({col}{R0}:{col}{TR-1})')
raw.cell(row=TR, column=18, value=f'=D{TR}-Q{TR}')
for cix in range(1, 19):
    c = raw.cell(row=TR, column=cix)
    c.font = F(10, True); c.fill = TOT_FILL; c.border = Border(left=thin, right=thin, top=med, bottom=med)
    if cix >= 4: c.number_format = N0; c.alignment = R
raw.cell(row=TR+2, column=1,
         value='검증 열 = 입고수량(양품+불량계) − 원자료 합계. 전 행 0 이어야 정상이다.').font = F(9, color='FF5A6672', it=True)
widths(raw, {'A':5,'B':15,'C':10,'D':11,'E':10,'F':10,'Q':12,'R':8})
for i in range(10): raw.column_dimensions[gcl(7+i)].width = 11
raw.freeze_panes = 'D5'

RAWCOL = {d: gcl(7+i) for i, d in enumerate(DEFECTS)}   # 유형 -> 원본데이터 열문자
LOT_R  = (R0, TR-1)

# =========================================================
# 요약
# =========================================================
s = wb.create_sheet('요약', 0)
title(s, 'U20 플라스틱 프리즘 — 입고검사 불량 분석', '2024-07-09 ~ 2026-01-23 · 8개 납품 로트 누적 · 모든 수치는 원본데이터 시트를 참조하는 수식이다.', 6)

s['A4'] = '누적 실적'; s['A4'].font = F(12, True, 'FF1F3A54')
kpi = [
    ('검사 로트 수', f'=COUNT(원본데이터!A{LOT_R[0]}:A{LOT_R[1]})', N0, False),
    ('검사 기간', '2024-07-09 ~ 2026-01-23', None, False),
    ('총 입고수량', f'=원본데이터!D{TR}', N0, True),
    ('총 양품수량', f'=원본데이터!E{TR}', N0, False),
    ('총 불량수량', f'=원본데이터!F{TR}', N0, True),
    ('누적 양품률', f'=원본데이터!E{TR}/원본데이터!D{TR}', P2, False),
    ('누적 불량률', f'=원본데이터!F{TR}/원본데이터!D{TR}', P2, True),
]
head(s, 5, ['항목','값'], h=20)
for i, (k, v, fmt, key) in enumerate(kpi):
    r = 6+i
    a = s.cell(row=r, column=1, value=k); a.font = F(10, key); a.border = BOX; a.alignment = L_
    b = s.cell(row=r, column=2, value=v); b.font = F(11, key); b.border = BOX; b.alignment = R
    if fmt: b.number_format = fmt
    if key: a.fill = KEY_FILL; b.fill = KEY_FILL

s['A15'] = '불량 유형 TOP 3'; s['A15'].font = F(12, True, 'FF1F3A54')
head(s, 16, ['순위','불량 유형','수량','입고 대비','불량 중 비중'], h=20)
order = sorted(DEFECTS, key=lambda d: -sum(L['counts'][DEFECTS.index(d)] for L in lots))
for i, d in enumerate(order[:3]):
    r = 17+i; col = RAWCOL[d]
    vals = [i+1, d, f'=원본데이터!{col}{TR}',
            f'=원본데이터!{col}{TR}/원본데이터!$D${TR}',
            f'=원본데이터!{col}{TR}/원본데이터!$F${TR}']
    for j, v in enumerate(vals):
        c = s.cell(row=r, column=1+j, value=v); c.font = F(10, i == 0); c.border = BOX
        c.alignment = C if j == 0 else (L_ if j == 1 else R)
        if j == 2: c.number_format = N0
        if j >= 3: c.number_format = P2
        if i == 0: c.fill = KEY_FILL

s['A21'] = '핵심 관찰'; s['A21'].font = F(12, True, 'FF1F3A54')
notes = [
 '1. 표면 얼룩은 만성 불량이 아니라 마지막 로트에서 터진 사건이다.',
 '   누적 6,088건 중 5,843건(96.0%)이 2026-01-23 (2차) 한 로트에 집중돼 있다. 이전 7개 로트 합계는 245건뿐이다.',
 '2. 같은 날 입고된 2026-01-23 (1차)와 (2차)의 불량률이 6.94% 대 44.92%로 극단적으로 다르다.',
 '   같은 날짜라도 생산 로트·금형·시기가 다르다는 뜻이므로, 두 건을 하나로 합산해 해석하면 안 된다.',
 '3. 공정은 한 번 안정화됐다가 다시 무너졌다. 37.60% → 24.78% → 2.38% → 1.57% → 0.58%까지 개선된 뒤',
 '   2026년에 6.94% → 44.92%로 급등했다. 개선 방법을 이미 아는 공정이므로 원인 추적이 가능하다.',
 '4. 흰색 이물도 2024-11-14에 1,036건(20.56%) 스파이크가 있었다. 단발성 사건이 반복되는 패턴이다.',
 '5. 2026-01-23 두 로트 모두에서 흰색 이물·스크래치가 동반 증가했다. 특정 유형이 아니라 공정 전반의 악화 신호다.',
 '',
 '※ 검사 기준서(7종)와 본 자료의 분류(10종)는 항목이 다르다. 기준서의 정식 항목인 「내부 가스」가 본 자료에 없고,',
 '   본 자료의 「내부 백테」·「표면 얼룩」에 흡수돼 있을 가능성이 있다. 분류 체계 정합성 확인이 필요하다.',
]
for i, t in enumerate(notes):
    c = s.cell(row=22+i, column=1, value=t)
    c.font = F(9.5, t[:2] in ('1.','2.','3.','4.','5.'), 'FF5A6672' if t.startswith('   ') or t.startswith('※') else 'FF1F2328')
    c.alignment = L_
widths(s, {'A':46,'B':22,'C':13,'D':13,'E':14})

# =========================================================
# 로트별 분석 (과거 → 현재)
# =========================================================
lt = wb.create_sheet('로트별분석', 1)
title(lt, '로트별 분석 — 과거부터 현재까지', '납품 순서대로 정렬. 불량률 추이와 로트별 최다 불량 유형.', 9)
head(lt, 4, ['No','납품일','연월','입고수량','양품','불량','양품률','불량률','최다 불량 유형','해당 수량','불량 중 비중'])
for i in range(len(lots)):
    r, sr = 5+i, R0+i
    vals = [f'=원본데이터!A{sr}', f'=원본데이터!B{sr}', f'=원본데이터!C{sr}',
            f'=원본데이터!D{sr}', f'=원본데이터!E{sr}', f'=원본데이터!F{sr}',
            f'=E{r}/D{r}', f'=F{r}/D{r}',
            f'=INDEX(원본데이터!$G$4:$P$4,MATCH(MAX(원본데이터!$G{sr}:$P{sr}),원본데이터!$G{sr}:$P{sr},0))',
            f'=MAX(원본데이터!$G{sr}:$P{sr})', f'=J{r}/F{r}']
    for j, v in enumerate(vals):
        c = lt.cell(row=r, column=1+j, value=v); c.font = F(10); c.border = BOX
        c.alignment = C if j in (0,2) else (L_ if j in (1,8) else R)
        if j in (3,4,5,9): c.number_format = N0
        if j in (6,7,10):  c.number_format = P2
tr2 = 5+len(lots)
lt.cell(row=tr2, column=1, value='합계'); lt.merge_cells(start_row=tr2, start_column=1, end_row=tr2, end_column=3)
for cix, f_ in [(4,f'=원본데이터!D{TR}'),(5,f'=원본데이터!E{TR}'),(6,f'=원본데이터!F{TR}'),
                (7,f'=E{tr2}/D{tr2}'),(8,f'=F{tr2}/D{tr2}')]:
    lt.cell(row=tr2, column=cix, value=f_)
for cix in range(1, 12):
    c = lt.cell(row=tr2, column=cix); c.font = F(10, True); c.fill = TOT_FILL
    c.border = Border(left=thin, right=thin, top=med, bottom=med); c.alignment = R if cix >= 4 else C
    if cix in (4,5,6): c.number_format = N0
    if cix in (7,8):   c.number_format = P2
widths(lt, {'A':5,'B':16,'C':10,'D':11,'E':10,'F':10,'G':10,'H':10,'I':15,'J':11,'K':13})

ch = LineChart(); ch.title = '로트별 불량률 추이'; ch.height, ch.width = 8, 19
ch.y_axis.title = '불량률'; ch.y_axis.numFmt = '0%'; ch.x_axis.title = '납품 로트'
ch.add_data(Reference(lt, min_col=8, min_row=4, max_row=4+len(lots)), titles_from_data=True)
ch.set_categories(Reference(lt, min_col=2, min_row=5, max_row=4+len(lots)))
lt.add_chart(ch, 'A16')

# =========================================================
# 월별 집계
# =========================================================
mo = wb.create_sheet('월별집계', 2)
title(mo, '월별 집계', '같은 달에 2개 로트가 있는 2026-01은 합산했다. SUMIFS로 원본데이터를 집계한다.', 16)
head(mo, 4, ['연월','로트 수','입고수량','양품','불량','불량률'] + DEFECTS)
for i, m in enumerate(months):
    r = 5+i
    mo.cell(row=r, column=1, value=m).alignment = C
    mo.cell(row=r, column=2, value=f'=COUNTIF(원본데이터!$C${R0}:$C${TR-1},$A{r})')
    for k, col in enumerate(['D','E','F']):
        mo.cell(row=r, column=3+k,
                value=f'=SUMIFS(원본데이터!${col}${R0}:${col}${TR-1},원본데이터!$C${R0}:$C${TR-1},$A{r})')
    mo.cell(row=r, column=6, value=f'=E{r}/C{r}')
    for k, d in enumerate(DEFECTS):
        col = RAWCOL[d]
        mo.cell(row=r, column=7+k,
                value=f'=SUMIFS(원본데이터!${col}${R0}:${col}${TR-1},원본데이터!$C${R0}:$C${TR-1},$A{r})')
    for cix in range(1, 17):
        c = mo.cell(row=r, column=cix); c.font = F(10); c.border = BOX
        c.alignment = C if cix == 1 else R
        if cix != 6 and cix != 1: c.number_format = N0
        if cix == 6: c.number_format = P2
tr3 = 5+len(months)
mo.cell(row=tr3, column=1, value='합계').alignment = C
for cix in range(2, 17):
    col = gcl(cix)
    mo.cell(row=tr3, column=cix, value=f'=SUM({col}5:{col}{tr3-1})')
mo.cell(row=tr3, column=6, value=f'=E{tr3}/C{tr3}')
for cix in range(1, 17):
    c = mo.cell(row=tr3, column=cix); c.font = F(10, True); c.fill = TOT_FILL
    c.border = Border(left=thin, right=thin, top=med, bottom=med); c.alignment = C if cix == 1 else R
    if cix not in (1,6): c.number_format = N0
    if cix == 6: c.number_format = P2
widths(mo, {'A':11,'B':9,'C':11,'D':10,'E':10,'F':10})
for i in range(10): mo.column_dimensions[gcl(7+i)].width = 11
mo.freeze_panes = 'C5'

# =========================================================
# 불량유형별 (파레토)
# =========================================================
dt = wb.create_sheet('불량유형별', 3)
title(dt, '불량 유형별 분석 (파레토)', '누적 수량 내림차순. 발생 로트 수는 해당 유형이 1건 이상 나온 로트의 개수다.', 8)
head(dt, 4, ['순위','불량 유형','누적 수량','입고 대비','불량 중 비중','누적 비중','발생 로트 수','최다 발생 로트','최다 수량'])
for i, d in enumerate(order):
    r = 5+i; col = RAWCOL[d]
    vals = [i+1, d, f'=원본데이터!{col}{TR}',
            f'=C{r}/원본데이터!$D${TR}', f'=C{r}/원본데이터!$F${TR}',
            f'=SUM($C$5:C{r})/원본데이터!$F${TR}',
            f'=COUNTIF(원본데이터!${col}${R0}:${col}${TR-1},">0")',
            f'=INDEX(원본데이터!$B${R0}:$B${TR-1},MATCH(MAX(원본데이터!${col}${R0}:${col}${TR-1}),원본데이터!${col}${R0}:${col}${TR-1},0))',
            f'=MAX(원본데이터!${col}${R0}:${col}${TR-1})']
    for j, v in enumerate(vals):
        c = dt.cell(row=r, column=1+j, value=v); c.font = F(10, i < 2); c.border = BOX
        c.alignment = C if j in (0,6) else (L_ if j in (1,7) else R)
        if j in (2,8): c.number_format = N0
        if j in (3,4,5): c.number_format = P2
        if i < 2: c.fill = KEY_FILL
tr4 = 5+len(order)
dt.cell(row=tr4, column=1, value='합계'); dt.merge_cells(start_row=tr4, start_column=1, end_row=tr4, end_column=2)
dt.cell(row=tr4, column=3, value=f'=SUM(C5:C{tr4-1})')
dt.cell(row=tr4, column=4, value=f'=C{tr4}/원본데이터!$D${TR}')
dt.cell(row=tr4, column=5, value=f'=C{tr4}/원본데이터!$F${TR}')
for cix in range(1, 10):
    c = dt.cell(row=tr4, column=cix); c.font = F(10, True); c.fill = TOT_FILL
    c.border = Border(left=thin, right=thin, top=med, bottom=med); c.alignment = C if cix <= 2 else R
    if cix == 3: c.number_format = N0
    if cix in (4,5): c.number_format = P2
widths(dt, {'A':6,'B':16,'C':12,'D':11,'E':13,'F':11,'G':13,'H':17,'I':11})

bc = BarChart(); bc.type = 'col'; bc.title = '불량 유형별 누적 수량'; bc.height, bc.width = 9, 19
bc.y_axis.title = '수량 (EA)'
bc.add_data(Reference(dt, min_col=3, min_row=4, max_row=4+len(order)), titles_from_data=True)
bc.set_categories(Reference(dt, min_col=2, min_row=5, max_row=4+len(order)))
dt.add_chart(bc, 'A18')

# =========================================================
# 유형별 추이 (로트 × 유형)
# =========================================================
tt = wb.create_sheet('유형별추이', 4)
title(tt, '유형별 발생 추이 — 로트 × 불량 유형', '위 표는 수량, 아래 표는 각 로트 입고수량 대비 발생률.', 12)
head(tt, 4, ['No','납품일','입고수량'] + order)
for i in range(len(lots)):
    r, sr = 5+i, R0+i
    tt.cell(row=r, column=1, value=f'=원본데이터!A{sr}')
    tt.cell(row=r, column=2, value=f'=원본데이터!B{sr}')
    tt.cell(row=r, column=3, value=f'=원본데이터!D{sr}')
    for j, d in enumerate(order):
        tt.cell(row=r, column=4+j, value=f'=원본데이터!{RAWCOL[d]}{sr}')
    for cix in range(1, 14):
        c = tt.cell(row=r, column=cix); c.font = F(10); c.border = BOX
        c.alignment = C if cix == 1 else (L_ if cix == 2 else R)
        if cix >= 3: c.number_format = N0
tr5 = 5+len(lots)
tt.cell(row=tr5, column=1, value='합계'); tt.merge_cells(start_row=tr5, start_column=1, end_row=tr5, end_column=2)
for cix in range(3, 14):
    col = gcl(cix); tt.cell(row=tr5, column=cix, value=f'=SUM({col}5:{col}{tr5-1})')
for cix in range(1, 14):
    c = tt.cell(row=tr5, column=cix); c.font = F(10, True); c.fill = TOT_FILL
    c.border = Border(left=thin, right=thin, top=med, bottom=med); c.alignment = R if cix >= 3 else C
    if cix >= 3: c.number_format = N0

B0 = tr5 + 3
tt.cell(row=B0-1, column=1, value='입고수량 대비 발생률').font = F(12, True, 'FF1F3A54')
head(tt, B0, ['No','납품일','불량률'] + order)
for i in range(len(lots)):
    r, sr, ur = B0+1+i, R0+i, 5+i
    tt.cell(row=r, column=1, value=f'=원본데이터!A{sr}')
    tt.cell(row=r, column=2, value=f'=원본데이터!B{sr}')
    tt.cell(row=r, column=3, value=f'=원본데이터!F{sr}/원본데이터!D{sr}')
    for j in range(len(order)):
        col = gcl(4+j)
        tt.cell(row=r, column=4+j, value=f'={col}{ur}/$C{ur}')
    for cix in range(1, 14):
        c = tt.cell(row=r, column=cix); c.font = F(10); c.border = BOX
        c.alignment = C if cix == 1 else (L_ if cix == 2 else R)
        if cix >= 3: c.number_format = P2
widths(tt, {'A':5,'B':16,'C':11})
for i in range(10): tt.column_dimensions[gcl(4+i)].width = 12
tt.freeze_panes = 'C5'

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = False

wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print('saved:', OUT)
print('lots:', len(lots), 'total:', sum(L['total'] for L in lots),
      'good:', sum(L['good'] for L in lots), 'defect:', sum(sum(L['counts']) for L in lots))
